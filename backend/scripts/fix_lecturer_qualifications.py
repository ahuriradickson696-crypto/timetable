"""Read 04_Lecturers sheet and populate LecturerModule qualification links."""
import sys, os
sys.path.insert(0, os.path.abspath(os.path.dirname(os.path.dirname(__file__))))

from openpyxl import load_workbook
from app.db.session import SessionLocal
from app.models.lecturer import Lecturer, LecturerModule
from app.models.course_unit import CourseUnit

WORKBOOK = r"C:\Users\MAS GLOBAL TECH\aviu-timetable\data\AVIU_Timetable_Master_Data_Workbook.xlsx"

db = SessionLocal()
try:
    print("Loading workbook...")
    wb = load_workbook(WORKBOOK, data_only=True)
    ws = wb["04_Lecturers"]

    # Find header row
    hdr_row = None
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=15, values_only=True), 1):
        for cell in row:
            if cell is None:
                continue
            if str(cell).strip().lower() == "lecturer_id":
                hdr_row = i
                break
            break
        if hdr_row:
            break

    print(f"Header row: {hdr_row}")
    header = [str(c).strip() if c else "" for c in next(ws.iter_rows(min_row=hdr_row, max_row=hdr_row, values_only=True))]
    print(f"Columns: {header}")

    # Clear existing links (safe to re-run)
    db.query(LecturerModule).delete()
    db.commit()

    added = 0
    skipped = 0
    missing_lecturers = 0
    for row in ws.iter_rows(min_row=hdr_row + 1, values_only=True):
        if not any(row):
            continue
        d = dict(zip(header, row))
        lid_raw = d.get("Lecturer_ID")
        if not lid_raw:
            continue
        lid = str(lid_raw).strip()
        if not db.get(Lecturer, lid):
            missing_lecturers += 1
            continue
        can_teach = d.get("Can_Teach_Modules") or ""
        modules = [m.strip() for m in str(can_teach).split(",") if m.strip()]
        for m in modules:
            if db.get(CourseUnit, m):
                db.add(LecturerModule(lecturer_id=lid, course_unit_code=m))
                added += 1
            else:
                skipped += 1
    db.commit()

    print()
    print("=" * 60)
    print("LECTURER QUALIFICATIONS IMPORT")
    print("=" * 60)
    print(f"  LecturerModule links added:  {added}")
    print(f"  Skipped (unknown modules):   {skipped}")
    print(f"  Skipped (lecturer missing):  {missing_lecturers}")
    print(f"  Total links now:             {db.query(LecturerModule).count()}")
    print("=" * 60)
finally:
    db.close()