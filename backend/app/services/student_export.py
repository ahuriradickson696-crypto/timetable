"""Professional student timetable exports - PDF, Excel, iCal.

Follows international university timetable format:
- Clean typography
- Full student details header
- Weekly grid with time bands
- Footer with reference and generation info
"""
import io
from typing import List, Dict, Any
from datetime import datetime, timedelta

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.platypus import (
    SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, PageBreak
)


AVIU_PURPLE = "4A148C"          # Deeper purple (AVIU brand)
AVIU_PURPLE_LIGHT = "F3E5F5"
AVIU_ACCENT = "6B21A8"
WEEKDAY_ORDER = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]
WEEKDAY_SHORT = {"Monday": "MON", "Tuesday": "TUE", "Wednesday": "WED",
                 "Thursday": "THU", "Friday": "FRI", "Saturday": "SAT", "Sunday": "SUN"}


def _cell_to_dict(c) -> Dict[str, Any]:
    if isinstance(c, dict):
        return c
    return {
        "day": getattr(c, "day", "") or "",
        "start_time": getattr(c, "start_time", "") or "",
        "end_time": getattr(c, "end_time", "") or "",
        "module_code": getattr(c, "module_code", "") or "",
        "module_name": getattr(c, "module_name", "") or "",
        "room_code": getattr(c, "room_code", "") or "",
        "lecturer_name": getattr(c, "lecturer_name", "") or "",
        "cohort_id": getattr(c, "cohort_id", "") or "",
    }


def _time_band(start: str, end: str) -> str:
    """Format time band nicely: 08:00 - 10:00"""
    return start + " - " + end


# ============================================================
# PDF -- Professional University Format
# ============================================================
def build_student_pdf(title, student_info, cells):
    data = [_cell_to_dict(c) for c in cells]
    data = [c for c in data if c["module_code"]]

    days_used = [d for d in WEEKDAY_ORDER if any(c["day"] == d for c in data)]
    if not days_used:
        days_used = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday"]

    time_slots = sorted({(c["start_time"], c["end_time"]) for c in data})

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=landscape(A4),
        topMargin=12 * mm,
        bottomMargin=12 * mm,
        leftMargin=12 * mm,
        rightMargin=12 * mm,
        title=title,
        author="AVIU Academic Registrar",
    )

    styles = getSampleStyleSheet()

    # Custom styles
    uni_name = ParagraphStyle(
        "UniName", parent=styles["Heading1"],
        textColor=colors.HexColor("#" + AVIU_PURPLE),
        fontSize=18, alignment=TA_CENTER, spaceAfter=2,
        fontName="Helvetica-Bold",
    )
    campus = ParagraphStyle(
        "Campus", parent=styles["Normal"],
        textColor=colors.HexColor("#666666"),
        fontSize=9, alignment=TA_CENTER, spaceAfter=8,
        fontName="Helvetica-Oblique",
    )
    doc_title = ParagraphStyle(
        "DocTitle", parent=styles["Heading2"],
        textColor=colors.HexColor("#" + AVIU_ACCENT),
        fontSize=13, alignment=TA_LEFT, spaceAfter=2,
        fontName="Helvetica-Bold",
    )
    doc_sub = ParagraphStyle(
        "DocSub", parent=styles["Normal"],
        textColor=colors.HexColor("#333333"),
        fontSize=9, alignment=TA_LEFT,
    )
    footer_style = ParagraphStyle(
        "Footer", parent=styles["Normal"],
        textColor=colors.HexColor("#666666"),
        fontSize=8, alignment=TA_CENTER,
    )

    story = []

    # ================== HEADER ==================
    story.append(Paragraph("AVANCE INTERNATIONAL UNIVERSITY", uni_name))
    story.append(Paragraph("Nabweru Campus, Kampala, Uganda", campus))

    # Horizontal rule
    rule = Table([[""]], colWidths=[273 * mm], rowHeights=[1])
    rule.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#" + AVIU_PURPLE)),
    ]))
    story.append(rule)
    story.append(Spacer(1, 4 * mm))

    # ================== STUDENT INFO ==================
    # Left: student details, Right: semester info
    info_left = [
        ["Name:", student_info.get("Name", "-")],
        ["Student No:", student_info.get("Student No.", "-")],
        ["Programme:", student_info.get("Programme", "-")],
    ]
    info_right = [
        ["Academic Year:", student_info.get("Academic Year", "-")],
        ["Semester:", student_info.get("Semester", "-")],
        ["Mode:", student_info.get("Mode", "-")],
    ]
    if "Combination" in student_info and student_info["Combination"]:
        info_left.append(["Combination:", student_info["Combination"]])
    if "Year" in student_info:
        info_right.append(["Year of Study:", student_info["Year"]])

    left_table = Table(info_left, colWidths=[35 * mm, 100 * mm])
    left_table.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
        ("FONTNAME", (1, 0), (1, -1), "Helvetica"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("TEXTCOLOR", (0, 0), (0, -1), colors.HexColor("#" + AVIU_ACCENT)),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
    ]))

    right_table = Table(info_right, colWidths=[35 * mm, 100 * mm])
    right_table.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
        ("FONTNAME", (1, 0), (1, -1), "Helvetica"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("TEXTCOLOR", (0, 0), (0, -1), colors.HexColor("#" + AVIU_ACCENT)),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
    ]))

    info_layout = Table([[left_table, right_table]], colWidths=[140 * mm, 133 * mm])
    info_layout.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 0),
        ("RIGHTPADDING", (0, 0), (-1, -1), 0),
    ]))
    story.append(info_layout)
    story.append(Spacer(1, 6 * mm))

    # ================== TIMETABLE GRID ==================
    if not data:
        story.append(Paragraph("No classes scheduled for this student.", styles["Normal"]))
    else:
        # Build grid
        header = ["TIME"] + [WEEKDAY_SHORT.get(d, d.upper()) for d in days_used]
        rows = [header]

        for start, end in time_slots:
            row = [_time_band(start, end)]
            for day in days_used:
                match = None
                for c in data:
                    if c["day"] == day and c["start_time"] == start and c["end_time"] == end:
                        match = c
                        break
                if match:
                    lines = [match["module_code"]]
                    if match["module_name"]:
                        lines.append(match["module_name"])
                    if match["room_code"]:
                        lines.append("Room " + match["room_code"])
                    if match["lecturer_name"]:
                        lines.append(match["lecturer_name"])
                    row.append("\n".join(lines))
                else:
                    row.append("")
            rows.append(row)

        # Column widths (landscape A4 = 297mm - margins = ~273mm usable)
        time_col_w = 28 * mm
        day_col_w = (273 * mm - time_col_w) / len(days_used)
        col_widths = [time_col_w] + [day_col_w] * len(days_used)

        tbl = Table(rows, colWidths=col_widths, repeatRows=1)
        tbl.setStyle(TableStyle([
            # Header
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + AVIU_PURPLE)),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, 0), 10),
            ("ALIGN", (0, 0), (-1, 0), "CENTER"),
            ("VALIGN", (0, 0), (-1, 0), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, 0), 6),
            ("BOTTOMPADDING", (0, 0), (-1, 0), 6),

            # Time column
            ("BACKGROUND", (0, 1), (0, -1), colors.HexColor("#" + AVIU_PURPLE_LIGHT)),
            ("FONTNAME", (0, 1), (0, -1), "Helvetica-Bold"),
            ("FONTSIZE", (0, 1), (0, -1), 8),
            ("TEXTCOLOR", (0, 1), (0, -1), colors.HexColor("#" + AVIU_PURPLE)),
            ("ALIGN", (0, 1), (0, -1), "CENTER"),
            ("VALIGN", (0, 1), (0, -1), "MIDDLE"),

            # Body cells
            ("FONTSIZE", (1, 1), (-1, -1), 7.5),
            ("VALIGN", (1, 1), (-1, -1), "TOP"),
            ("LEFTPADDING", (1, 1), (-1, -1), 5),
            ("RIGHTPADDING", (1, 1), (-1, -1), 5),
            ("TOPPADDING", (1, 1), (-1, -1), 6),
            ("BOTTOMPADDING", (1, 1), (-1, -1), 6),

            # Grid lines
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#DDDDDD")),
            ("BOX", (0, 0), (-1, -1), 1, colors.HexColor("#" + AVIU_PURPLE)),
        ]))
        story.append(tbl)

    # ================== FOOTER ==================
    story.append(Spacer(1, 8 * mm))

    # Reference number
    ref = "TIM/" + datetime.now().strftime("%Y/%m") + "/" + str(datetime.now().strftime("%H%M%S"))
    gen_time = datetime.now().strftime("%d %B %Y at %H:%M")

    footer_text = (
        "Prepared by: Office of the Academic Registrar &nbsp;&nbsp;|&nbsp;&nbsp; "
        "Reference: " + ref + " &nbsp;&nbsp;|&nbsp;&nbsp; "
        "Generated: " + gen_time
    )
    story.append(Paragraph(footer_text, footer_style))

    story.append(Spacer(1, 3 * mm))
    story.append(Paragraph(
        "This document is auto-generated by the AVIU Timetabling System. "
        "Any unauthorized modification is not valid.",
        ParagraphStyle("Disclaimer", parent=footer_style, fontSize=7,
                       textColor=colors.HexColor("#999999")),
    ))

    doc.build(story)
    buf.seek(0)
    return buf.read()


# ============================================================
# EXCEL -- Professional Format
# ============================================================
def build_student_excel(title, student_info, cells):
    data = [_cell_to_dict(c) for c in cells]
    data = [c for c in data if c["module_code"]]

    days_used = [d for d in WEEKDAY_ORDER if any(c["day"] == d for c in data)]
    if not days_used:
        days_used = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday"]
    time_slots = sorted({(c["start_time"], c["end_time"]) for c in data})

    wb = Workbook()
    ws = wb.active
    ws.title = "Timetable"

    purple_fill = PatternFill(start_color=AVIU_PURPLE, end_color=AVIU_PURPLE, fill_type="solid")
    light_fill = PatternFill(start_color=AVIU_PURPLE_LIGHT, end_color=AVIU_PURPLE_LIGHT, fill_type="solid")
    white_font = Font(bold=True, color="FFFFFF")
    bold_font = Font(bold=True, size=11, color=AVIU_ACCENT)

    total_cols = 1 + len(days_used)

    # Row 1: University name
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_cols)
    t = ws.cell(row=1, column=1, value="AVANCE INTERNATIONAL UNIVERSITY")
    t.font = Font(bold=True, size=16, color="FFFFFF")
    t.fill = purple_fill
    t.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 32

    # Row 2: Campus
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=total_cols)
    s = ws.cell(row=2, column=1, value="Nabweru Campus, Kampala - Student Timetable")
    s.font = Font(italic=True, size=10, color=AVIU_ACCENT)
    s.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 20

    # Row 4+: Student info
    info_row = 4
    pairs = list(student_info.items())
    for i in range(0, len(pairs), 2):
        k1, v1 = pairs[i]
        c1 = ws.cell(row=info_row, column=1, value=k1)
        c1.font = bold_font
        c1.fill = light_fill
        ws.cell(row=info_row, column=2, value=v1)
        if i + 1 < len(pairs):
            k2, v2 = pairs[i + 1]
            if total_cols >= 4:
                c3 = ws.cell(row=info_row, column=3, value=k2)
                c3.font = bold_font
                c3.fill = light_fill
                ws.cell(row=info_row, column=4, value=v2)
        info_row += 1

    # Timetable header
    header_row = info_row + 1
    ws.cell(row=header_row, column=1, value="TIME")
    for col_idx, day in enumerate(days_used, start=2):
        ws.cell(row=header_row, column=col_idx, value=WEEKDAY_SHORT.get(day, day.upper()))

    for col_idx in range(1, total_cols + 1):
        c = ws.cell(row=header_row, column=col_idx)
        c.font = white_font
        c.fill = purple_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[header_row].height = 24

    thin = Side(border_style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    row = header_row + 1
    for start, end in time_slots:
        tc = ws.cell(row=row, column=1, value=_time_band(start, end))
        tc.font = Font(bold=True, size=9, color=AVIU_ACCENT)
        tc.fill = light_fill
        tc.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        tc.border = border
        for col_idx, day in enumerate(days_used, start=2):
            match = None
            for c in data:
                if c["day"] == day and c["start_time"] == start and c["end_time"] == end:
                    match = c
                    break
            if match:
                text = match["module_code"] + "\n" + match["module_name"]
                if match["room_code"]:
                    text += "\nRoom: " + match["room_code"]
                if match["lecturer_name"]:
                    text += "\n" + match["lecturer_name"]
            else:
                text = ""
            cell = ws.cell(row=row, column=col_idx, value=text)
            cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            cell.border = border
            if match:
                cell.fill = light_fill
        ws.row_dimensions[row].height = 75
        row += 1

    ws.column_dimensions["A"].width = 14
    for i in range(2, total_cols + 1):
        ws.column_dimensions[get_column_letter(i)].width = 30

    row += 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=total_cols)
    f = ws.cell(row=row, column=1,
                value="Prepared by the Office of the Academic Registrar | Generated by AVIU Timetabling System on " + datetime.now().strftime("%d %b %Y at %H:%M"))
    f.font = Font(italic=True, size=8, color="666666")
    f.alignment = Alignment(horizontal="center")

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


# ============================================================
# iCal -- Simple valid format
# ============================================================
def build_student_ical(title, cells):
    data = [_cell_to_dict(c) for c in cells]
    data = [c for c in data if c["module_code"]]

    lines = [
        "BEGIN:VCALENDAR",
        "VERSION:2.0",
        "PRODID:-//AVIU Timetable//EN",
        "CALSCALE:GREGORIAN",
        "METHOD:PUBLISH",
        "X-WR-CALNAME:" + title,
        "X-WR-TIMEZONE:Africa/Kampala",
    ]

    today = datetime.now()
    days_until_monday = (7 - today.weekday()) % 7
    next_monday = today + timedelta(days=days_until_monday)

    for i, c in enumerate(data):
        try:
            day_offset = WEEKDAY_ORDER.index(c["day"])
        except ValueError:
            continue
        start_dt = next_monday + timedelta(days=day_offset)
        try:
            sh, sm = map(int, c["start_time"].split(":"))
            eh, em = map(int, c["end_time"].split(":"))
        except ValueError:
            continue
        start = start_dt.replace(hour=sh, minute=sm, second=0, microsecond=0)
        end = start_dt.replace(hour=eh, minute=em, second=0, microsecond=0)

        summary = c["module_code"]
        if c["module_name"]:
            summary = summary + " - " + c["module_name"]
        desc_parts = []
        if c["lecturer_name"]:
            desc_parts.append("Lecturer: " + c["lecturer_name"])
        if c["cohort_id"]:
            desc_parts.append("Cohort: " + c["cohort_id"])

        uid = "aviu-" + str(i) + "-" + start.strftime("%Y%m%d%H%M") + "@aviu.ac.ug"
        dtstamp = datetime.utcnow().strftime("%Y%m%dT%H%M%SZ")
        dtstart = start.strftime("%Y%m%dT%H%M%S")
        dtend = end.strftime("%Y%m%dT%H%M%S")

        lines.append("BEGIN:VEVENT")
        lines.append("UID:" + uid)
        lines.append("DTSTAMP:" + dtstamp)
        lines.append("DTSTART:" + dtstart)
        lines.append("DTEND:" + dtend)
        lines.append("SUMMARY:" + summary)
        lines.append("LOCATION:" + c["room_code"])
        lines.append("DESCRIPTION:" + " | ".join(desc_parts))
        lines.append("END:VEVENT")

    lines.append("END:VCALENDAR")
    return ("\r\n".join(lines)).encode("utf-8")
