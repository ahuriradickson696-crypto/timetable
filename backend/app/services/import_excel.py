"""Smart Excel importer for the AVIU Master Data Workbook."""
import os
from typing import Dict, Any, List, Optional
from datetime import datetime, date

from openpyxl import load_workbook
from sqlalchemy.orm import Session
from sqlalchemy import select

from app.models.programme import Programme
from app.models.cohort import Cohort
from app.models.course_unit import CourseUnit
from app.models.lecturer import Lecturer
from app.models.room import Room
from app.models.time_period import TimePeriod
from app.models.activity import Activity
from app.core.logging import logger


def _clean(value: Any) -> Optional[str]:
    if value is None:
        return None
    s = str(value).strip()
    return s if s else None


def _to_int(value: Any, default: Optional[int] = None) -> Optional[int]:
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return default


def _to_date(value: Any) -> Optional[date]:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    s = str(value).strip()
    for fmt in ("%d-%m-%Y", "%Y-%m-%d", "%d/%m/%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def _find_header_row(ws, marker: str, max_scan: int = 15) -> Optional[int]:
    """Find the row where the FIRST cell EXACTLY equals marker (case-insensitive)."""
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=max_scan, values_only=True), start=1):
        for cell in row:
            if cell is None:
                continue
            if str(cell).strip().lower() == marker.lower():
                return i
            break  # only check the first non-empty cell per row
    return None


def _read_sheet(ws, marker: str) -> List[Dict[str, Any]]:
    hdr_row = _find_header_row(ws, marker)
    if hdr_row is None:
        return []
    header_row = next(ws.iter_rows(min_row=hdr_row, max_row=hdr_row, values_only=True))
    header = [(_clean(c) or "") for c in header_row]
    rows = []
    for row in ws.iter_rows(min_row=hdr_row + 1, values_only=True):
        if not any(row):
            continue
        d = {}
        for k, v in zip(header, row):
            if k:
                d[k] = v
        first_key = header[0] if header else None
        if first_key and d.get(first_key):
            rows.append(d)
    return rows


def import_from_excel(db: Session, file_path: str) -> Dict[str, Any]:
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"Workbook not found: {file_path}")

    logger.info(f"Loading workbook: {file_path}")
    wb = load_workbook(file_path, data_only=True)

    stats = {
        "programmes": {"inserted": 0, "updated": 0},
        "cohorts": {"inserted": 0, "updated": 0},
        "course_units": {"inserted": 0, "updated": 0},
        "lecturers": {"inserted": 0, "updated": 0},
        "rooms": {"inserted": 0, "updated": 0},
        "time_periods": {"inserted": 0, "updated": 0},
        "activities": {"inserted": 0, "updated": 0},
    }
    warnings: List[str] = []

    # 01_Programmes
    if "01_Programmes" in wb.sheetnames:
        rows = _read_sheet(wb["01_Programmes"], "Programme_Code")
        for r in rows:
            code = _clean(r.get("Programme_Code"))
            if not code:
                continue
            existing = db.get(Programme, code)
            data = dict(
                name=_clean(r.get("Programme_Name")) or code,
                level=_clean(r.get("Level")) or "Bachelor",
                school_faculty=_clean(r.get("School_Faculty")) or "Unknown",
                tier_discipline=_clean(r.get("Tier_Discipline")),
                duration_years=_to_int(r.get("Duration_Years"), 3) or 3,
                last_accreditation=_to_date(r.get("Last_Accreditation")),
                review_year=_to_int(r.get("Review_Year")),
                active=(_clean(r.get("Active")) or "Yes").lower() == "yes",
                notes=_clean(r.get("Notes")),
            )
            if existing:
                for k, v in data.items():
                    setattr(existing, k, v)
                stats["programmes"]["updated"] += 1
            else:
                db.add(Programme(code=code, **data))
                stats["programmes"]["inserted"] += 1
        db.commit()

    # 02_Cohorts
    if "02_Cohorts" in wb.sheetnames:
        rows = _read_sheet(wb["02_Cohorts"], "Cohort_ID")
        logger.info(f"Cohort rows read: {len(rows)}")
        for r in rows:
            cid = _clean(r.get("Cohort_ID"))
            if not cid:
                continue
            prog_code = _clean(r.get("Programme_Code"))
            if not prog_code or not db.get(Programme, prog_code):
                warnings.append(f"Cohort {cid}: programme {prog_code} missing")
                continue
            existing = db.get(Cohort, cid)
            data = dict(
                programme_code=prog_code,
                year_level=_to_int(r.get("Year_Level"), 1) or 1,
                mode=_clean(r.get("Mode")) or "Day",
                intake=_clean(r.get("Intake")) or "Aug",
                academic_year=_clean(r.get("Academic_Year")) or "2026/2027",
                semester=_to_int(r.get("Semester"), 1) or 1,
                estimated_students=_to_int(r.get("Estimated_Students"), 0) or 0,
                status=_clean(r.get("Status")) or "Active",
            )
            if existing:
                for k, v in data.items():
                    setattr(existing, k, v)
                stats["cohorts"]["updated"] += 1
            else:
                db.add(Cohort(id=cid, **data))
                stats["cohorts"]["inserted"] += 1
        db.commit()

    # 03_Course_Units
    if "03_Course_Units" in wb.sheetnames:
        rows = _read_sheet(wb["03_Course_Units"], "Module_Code")
        for r in rows:
            code = _clean(r.get("Module_Code"))
            if not code:
                continue
            existing = db.get(CourseUnit, code)
            data = dict(
                name=_clean(r.get("Module_Name")) or code,
                credit_units=_to_int(r.get("Credit_Units"), 3) or 3,
                contact_hours_per_week=_to_int(r.get("Contact_Hours_Per_Week"), 3) or 3,
                normal_year=_to_int(r.get("Normal_Year"), 1) or 1,
                is_cross_cutting=(_clean(r.get("Is_Cross_Cutting")) or "No").lower() == "yes",
                preferred_room_type=_clean(r.get("Preferred_Room_Type")) or "Lecture",
                owning_school=_clean(r.get("Owning_School")),
                active=(_clean(r.get("Active")) or "Yes").lower() == "yes",
                notes=_clean(r.get("Notes")),
            )
            if existing:
                for k, v in data.items():
                    setattr(existing, k, v)
                stats["course_units"]["updated"] += 1
            else:
                db.add(CourseUnit(code=code, **data))
                stats["course_units"]["inserted"] += 1
        db.commit()

    # 04_Lecturers
    if "04_Lecturers" in wb.sheetnames:
        rows = _read_sheet(wb["04_Lecturers"], "Lecturer_ID")
        for r in rows:
            lid = _clean(r.get("Lecturer_ID"))
            if not lid:
                continue
            existing = db.get(Lecturer, lid)
            data = dict(
                full_name=_clean(r.get("Full_Name")) or lid,
                max_hours_per_week=_to_int(r.get("Max_Hours_Per_Week"), 16) or 16,
                preferred_mode=_clean(r.get("Preferred_Mode")) or "Both",
                unavailable_days=_clean(r.get("Unavailable_Days")),
                email=_clean(r.get("Email")),
                active=(_clean(r.get("Active")) or "Yes").lower() == "yes",
                notes=_clean(r.get("Notes")),
            )
            if existing:
                for k, v in data.items():
                    setattr(existing, k, v)
                stats["lecturers"]["updated"] += 1
            else:
                db.add(Lecturer(id=lid, **data))
                stats["lecturers"]["inserted"] += 1
        db.commit()

    # 05_Rooms
    if "05_Rooms" in wb.sheetnames:
        rows = _read_sheet(wb["05_Rooms"], "Room_Code")
        for r in rows:
            code = _clean(r.get("Room_Code"))
            if not code:
                continue
            existing = db.get(Room, code)
            data = dict(
                name=_clean(r.get("Room_Name")) or code,
                capacity=_to_int(r.get("Capacity"), 30) or 30,
                room_type=_clean(r.get("Room_Type")) or "Lecture",
                building=_clean(r.get("Building")),
                features=_clean(r.get("Features")),
                available_modes=_clean(r.get("Available_Modes")) or "Both",
                active=True,
                notes=_clean(r.get("Notes")),
            )
            if existing:
                for k, v in data.items():
                    setattr(existing, k, v)
                stats["rooms"]["updated"] += 1
            else:
                db.add(Room(code=code, **data))
                stats["rooms"]["inserted"] += 1
        db.commit()

    # 06_Time_Structure
    if "06_Time_Structure" in wb.sheetnames:
        rows = _read_sheet(wb["06_Time_Structure"], "Period_ID")
        for r in rows:
            pid = _clean(r.get("Period_ID"))
            day = _clean(r.get("Day"))
            if not pid or not day:
                continue
            mode = _clean(r.get("Mode")) or "Day"
            existing = db.scalar(select(TimePeriod).where(
                TimePeriod.period_id == pid,
                TimePeriod.day == day,
                TimePeriod.mode == mode,
            ))
            data = dict(
                period_id=pid,
                day=day,
                start_time=_clean(r.get("Start_Time")) or "",
                end_time=_clean(r.get("End_Time")) or "",
                mode=mode,
                is_break=(_clean(r.get("Is_Break")) or "No").lower() == "yes",
                order_index=_to_int(r.get("Order_Index"), 0) or 0,
            )
            if existing:
                for k, v in data.items():
                    setattr(existing, k, v)
                stats["time_periods"]["updated"] += 1
            else:
                db.add(TimePeriod(**data))
                stats["time_periods"]["inserted"] += 1
        db.commit()

    # 07_Activities
    if "07_Activities" in wb.sheetnames:
        rows = _read_sheet(wb["07_Activities"], "Activity_ID")
        for r in rows:
            aid = _clean(r.get("Activity_ID"))
            if not aid:
                continue
            cohort_raw = _clean(r.get("Cohort_ID")) or ""
            parts = [p.strip() for p in cohort_raw.replace("+", ",").split(",") if p.strip()]
            if not parts:
                warnings.append(f"Activity {aid}: no cohort")
                continue
            primary_cohort = parts[0]
            joint_ids = ",".join(parts[1:]) if len(parts) > 1 else None
            module_raw = _clean(r.get("Module_Code")) or ""
            if "/" in module_raw:
                module_code = module_raw.split("/")[0].strip()
            else:
                module_code = module_raw
            if not module_code:
                warnings.append(f"Activity {aid}: no module")
                continue
            lect_raw = _clean(r.get("Lecturer_ID")) or ""
            lect_id = lect_raw.split(",")[0].strip() if lect_raw else ""
            if not lect_id:
                continue
            existing = db.get(Activity, aid)
            data = dict(
                cohort_id=primary_cohort,
                module_code=module_code,
                lecturer_id=lect_id,
                duration_periods=_to_int(r.get("Duration_Periods"), 1) or 1,
                preferred_room_type=_clean(r.get("Preferred_Room_Type")) or "Lecture",
                students=_to_int(r.get("Students"), 0) or 0,
                split_allowed=(_clean(r.get("Split_Allowed")) or "No").lower() == "yes",
                joint_cohort_ids=joint_ids,
                semester=1,
                intake="Aug",
                academic_year="2026/2027",
                notes=_clean(r.get("Notes")),
            )
            if existing:
                for k, v in data.items():
                    setattr(existing, k, v)
                stats["activities"]["updated"] += 1
            else:
                db.add(Activity(id=aid, **data))
                stats["activities"]["inserted"] += 1
        db.commit()

    logger.info(f"Import complete: {stats}")
    return {"status": "ok", "stats": stats, "warnings": warnings}
